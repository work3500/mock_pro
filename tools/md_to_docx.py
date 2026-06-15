#!/usr/bin/env python3
"""Markdown -> DOCX 변환기 (KTNET AI 코딩도구 활용 가이드 전용 서식).

표준 라이브러리 + python-docx 만 사용한다. 이 문서가 사용하는 마크다운 구성
요소(제목/표/펜스 코드블록/인용구/목록/구분선/인라인 강조·코드·링크)를
깔끔하고 가독성 좋은 한글 워드 문서로 변환하도록 서식을 직접 제어한다.
"""

import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# 색상/폰트 팔레트
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1F, 0x38, 0x64)        # 제목/표 헤더
BLUE = RGBColor(0x2E, 0x53, 0x95)        # 절 제목
SLATE = RGBColor(0x44, 0x54, 0x6A)       # 소절 제목
GREY_TEXT = RGBColor(0x59, 0x59, 0x59)   # 보조 텍스트
BODY_BLACK = RGBColor(0x22, 0x22, 0x22)

BODY_FONT = "맑은 고딕"          # 본문 (한글 Windows 표준)
CODE_FONT = "D2Coding"           # 코드/다이어그램 (고정폭 CJK)
CODE_FONT_FALLBACK = "Consolas"  # 라틴 고정폭

NAVY_HEX = "1F3864"
HDR_FILL = "1F3864"
CODE_FILL = "F4F5F7"
QUOTE_FILL = "FFF8E6"
QUOTE_BAR = "E0A800"


# ---------------------------------------------------------------------------
# 저수준 XML 헬퍼
# ---------------------------------------------------------------------------
def _set_run_fonts(run, latin, cjk=None):
    cjk = cjk or latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)


def _shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _paragraph_borders(paragraph, edges, color="D5D8DD", size=6, space=6):
    """edges: dict like {'top': True, 'left': {'color':..,'sz':..}}."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        if edge not in edges:
            continue
        spec = edges[edge]
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(spec.get("sz", size) if isinstance(spec, dict) else size))
        el.set(qn("w:space"), str(spec.get("space", space) if isinstance(spec, dict) else space))
        el.set(qn("w:color"), spec.get("color", color) if isinstance(spec, dict) else color)
        pbdr.append(el)
    ppr.append(pbdr)


def _set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _set_cell_vertical_center(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcpr.append(va)


def _set_table_width_pct(table, pct=5000):
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "pct")
    tblw.set(qn("w:w"), str(pct))


def _set_table_borders(table, color="C7CCD4", size=4):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def _set_cell_margins(table, top=40, bottom=40, left=90, right=90):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


# ---------------------------------------------------------------------------
# 인라인 파싱 (**bold**, *italic*, `code`, [text](url))
# ---------------------------------------------------------------------------
INLINE_RE = re.compile(
    r"(`[^`]+`|\*\*.+?\*\*|\*[^*\n]+?\*|\[[^\]]+\]\([^)]*\))"
)


def add_inline(paragraph, text, base_size=10.5, base_color=BODY_BLACK,
               base_font=BODY_FONT):
    """인라인 마크업을 해석해 run 들을 추가한다."""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.size = Pt(base_size - 0.5)
            _set_run_fonts(run, CODE_FONT_FALLBACK, CODE_FONT)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        elif token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
            _set_run_fonts(run, base_font)
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
            _set_run_fonts(run, base_font)
        elif token.startswith("[") and "](" in token:
            label = token[1:token.index("]")]
            run = paragraph.add_run(label)
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
            _set_run_fonts(run, base_font)
        else:
            run = paragraph.add_run(token)
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
            _set_run_fonts(run, base_font)


# ---------------------------------------------------------------------------
# 블록 렌더링
# ---------------------------------------------------------------------------
def add_heading(doc, level, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if level == 2:
        pf.page_break_before = True
        pf.space_before = Pt(6)
        pf.space_after = Pt(10)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(17)
        run.font.color.rgb = NAVY
        _set_run_fonts(run, BODY_FONT)
        _paragraph_borders(p, {"bottom": {"color": NAVY_HEX, "sz": 18, "space": 4}})
    elif level == 3:
        pf.space_before = Pt(14)
        pf.space_after = Pt(5)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13.5)
        run.font.color.rgb = BLUE
        _set_run_fonts(run, BODY_FONT)
    elif level == 4:
        pf.space_before = Pt(10)
        pf.space_after = Pt(3)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = SLATE
        _set_run_fonts(run, BODY_FONT)
    else:
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = NAVY
        _set_run_fonts(run, BODY_FONT)
    return p


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.4
    # 소프트 줄바꿈(두 칸 공백) 처리
    segments = text.split("\n")
    for i, seg in enumerate(segments):
        if i > 0:
            p.add_run().add_break(WD_BREAK.LINE)
        add_inline(p, seg)
    return p


def add_list_item(doc, level, marker, text, ordered=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    base = Cm(0.75) + Cm(0.6) * level
    pf.left_indent = base
    pf.first_line_indent = Cm(-0.55)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.35
    tab_stops = pf.tab_stops
    tab_stops.add_tab_stop(base, WD_TAB_ALIGNMENT.LEFT)
    bullet = marker if ordered else ("–" if level > 0 else "•")
    run = p.add_run(f"{bullet}\t")
    run.font.size = Pt(10.5)
    run.font.color.rgb = BLUE if not ordered else BODY_BLACK
    _set_run_fonts(run, BODY_FONT)
    add_inline(p, text)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.1)
    pf.right_indent = Cm(0.1)
    _shade_paragraph(p, CODE_FILL)
    _paragraph_borders(
        p,
        {
            "top": {"color": "D5D8DD", "sz": 6, "space": 6},
            "bottom": {"color": "D5D8DD", "sz": 6, "space": 6},
            "left": {"color": "D5D8DD", "sz": 6, "space": 6},
            "right": {"color": "D5D8DD", "sz": 6, "space": 6},
        },
    )
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break(WD_BREAK.LINE)
        run = p.add_run(line if line else "")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
        _set_run_fonts(run, CODE_FONT_FALLBACK, CODE_FONT)
    return p


def add_blockquote(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.left_indent = Cm(0.35)
    pf.line_spacing = 1.35
    _shade_paragraph(p, QUOTE_FILL)
    _paragraph_borders(
        p,
        {
            "left": {"color": QUOTE_BAR, "sz": 24, "space": 8},
            "top": {"color": "EFE3B8", "sz": 4, "space": 4},
            "bottom": {"color": "EFE3B8", "sz": 4, "space": 4},
            "right": {"color": "EFE3B8", "sz": 4, "space": 4},
        },
    )
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break(WD_BREAK.LINE)
        add_inline(p, line, base_size=10, base_color=RGBColor(0x4A, 0x40, 0x20))
    return p


def add_table(doc, rows, aligns):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_width_pct(table, 5000)
    _set_table_borders(table)
    _set_cell_margins(table)
    for r_idx, row in enumerate(rows):
        is_header = r_idx == 0
        for c_idx in range(ncols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            _set_cell_vertical_center(cell)
            if is_header:
                _set_cell_shading(cell, HDR_FILL)
            # 셀 기본 문단
            cell_p = cell.paragraphs[0]
            cell_p.paragraph_format.space_after = Pt(1)
            cell_p.paragraph_format.space_before = Pt(1)
            cell_p.paragraph_format.line_spacing = 1.2
            align = aligns[c_idx] if c_idx < len(aligns) else "left"
            cell_p.alignment = {
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "left": WD_ALIGN_PARAGRAPH.LEFT,
            }[align]
            if is_header:
                # 헤더 텍스트: 흰색 볼드
                for tok in INLINE_RE.split(text):
                    if not tok:
                        continue
                    clean = tok
                    if clean.startswith("**") and clean.endswith("**"):
                        clean = clean[2:-2]
                    run = cell_p.add_run(clean)
                    run.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    _set_run_fonts(run, BODY_FONT)
            else:
                add_inline(cell_p, text, base_size=9.5)
    return table


# ---------------------------------------------------------------------------
# 마크다운 파서
# ---------------------------------------------------------------------------
def parse_table_separator(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    aligns = []
    ok = True
    for c in cells:
        if not re.fullmatch(r":?-{1,}:?", c):
            ok = False
            break
        if c.startswith(":") and c.endswith(":"):
            aligns.append("center")
        elif c.endswith(":"):
            aligns.append("right")
        else:
            aligns.append("left")
    return (aligns if ok else None)


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    setup_document(doc)

    i = 0
    n = len(lines)
    seen_title = False
    para_buf = []

    def flush_para():
        nonlocal para_buf
        if not para_buf:
            return
        # para_buf: list of (text, hard_break_after)
        out = ""
        for idx, (txt, hard) in enumerate(para_buf):
            out += txt
            if idx < len(para_buf) - 1:
                out += "\n" if hard else " "
        add_body_paragraph(doc, out)
        para_buf = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 펜스 코드블록
        if stripped.startswith("```"):
            flush_para()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 닫는 ```
            # 앞뒤 빈 줄 제거
            while code_lines and not code_lines[0].strip():
                code_lines.pop(0)
            while code_lines and not code_lines[-1].strip():
                code_lines.pop()
            add_code_block(doc, code_lines)
            continue

        # 표
        if stripped.startswith("|") and i + 1 < n:
            aligns = parse_table_separator(lines[i + 1])
            if aligns is not None:
                flush_para()
                header = split_table_row(lines[i])
                rows = [header]
                i += 2
                while i < n and lines[i].strip().startswith("|"):
                    rows.append(split_table_row(lines[i]))
                    i += 1
                add_table(doc, rows, aligns)
                continue

        # 제목
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_para()
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not seen_title:
                add_cover_title(doc, text)
                seen_title = True
            else:
                add_heading(doc, level, text)
            i += 1
            continue

        # 인용구
        if stripped.startswith(">"):
            flush_para()
            q_lines = []
            while i < n and lines[i].strip().startswith(">"):
                q_lines.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            add_blockquote(doc, q_lines)
            continue

        # 구분선
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush_para()
            i += 1
            continue

        # 목록 (순서 없음)
        m = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
        if m:
            flush_para()
            indent = len(m.group(1))
            level = indent // 2
            add_list_item(doc, level, "", m.group(3))
            i += 1
            continue

        # 목록 (순서 있음)
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            flush_para()
            indent = len(m.group(1))
            level = indent // 3
            add_list_item(doc, level, f"{m.group(2)}.", m.group(3), ordered=True)
            i += 1
            continue

        # 빈 줄 -> 문단 분리
        if not stripped:
            flush_para()
            i += 1
            continue

        # 일반 본문 (두 칸 공백 = 강제 줄바꿈)
        hard_break = line.endswith("  ")
        para_buf.append((line.strip(), hard_break))
        i += 1

    flush_para()
    doc.save(docx_path)
    print(f"saved: {docx_path}")


def add_cover_title(doc, title):
    # 상단 여백
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(60)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = NAVY
    _set_run_fonts(run, BODY_FONT)
    # 강조 라인
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(18)
    _paragraph_borders(rule, {"bottom": {"color": NAVY_HEX, "sz": 18, "space": 1}})
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(24)
    srun = sub.add_run("한국무역정보통신(KTNET) · AX추진실")
    srun.font.size = Pt(12)
    srun.font.color.rgb = GREY_TEXT
    _set_run_fonts(srun, BODY_FONT)


def setup_document(doc):
    # 기본(Normal) 스타일
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY_BLACK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    rfonts.set(qn("w:cs"), BODY_FONT)

    # 페이지: A4 + 여백
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)

    _add_footer_page_number(doc)


def _add_footer_page_number(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run.font.size = Pt(9)
    run.font.color.rgb = GREY_TEXT
    _set_run_fonts(run, BODY_FONT)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
